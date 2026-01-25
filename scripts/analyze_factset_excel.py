#!/usr/bin/env python
"""Analyze FactSet public companies Excel file."""

import openpyxl
from collections import Counter

print("Loading Excel file...")
wb = openpyxl.load_workbook(
    r'B:\github\py-sec-edgar\data\FF_PUBLIC_COMPANIES_EXPORT_GLOBAL_IDENTIFIERS_NEW ALL.xlsx',
    read_only=True
)
ws = wb.active

# Get headers
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
print(f"Columns: {len(headers)}")

# Find key column indices
col_map = {h: i for i, h in enumerate(headers) if h}
print("\nKey columns:")
for key in ['ISIN', 'CUSIP', 'SEDOL', 'EntityID', 'Perm.Sec.ID', 'FECo InfoFactsetid', 'FECo InfoPermid', 'Company Name', 'PrimaryEquityListing']:
    if key in col_map:
        print(f"  {key}: column {col_map[key]}")

# Analyze data
rows = 0
nvidia_rows = []
isin_count = 0
permid_count = 0
entity_id_count = 0
perm_sec_count = 0

isin_col = col_map.get('ISIN')
permid_col = col_map.get('FECo InfoPermid')
entity_id_col = col_map.get('EntityID')
perm_sec_col = col_map.get('Perm.Sec.ID')
name_col = col_map.get('Company Name')
cusip_col = col_map.get('CUSIP')
sedol_col = col_map.get('SEDOL')
primary_eq_col = col_map.get('PrimaryEquityListing')

for row in ws.iter_rows(min_row=2, values_only=True):
    rows += 1
    
    name = str(row[name_col] or '').strip() if name_col else ''
    isin = str(row[isin_col] or '').strip() if isin_col and row[isin_col] else ''
    permid = str(row[permid_col] or '').strip() if permid_col and row[permid_col] else ''
    entity_id = str(row[entity_id_col] or '').strip() if entity_id_col and row[entity_id_col] else ''
    perm_sec = str(row[perm_sec_col] or '').strip() if perm_sec_col and row[perm_sec_col] else ''
    cusip = str(row[cusip_col] or '').strip() if cusip_col and row[cusip_col] else ''
    sedol = str(row[sedol_col] or '').strip() if sedol_col and row[sedol_col] else ''
    primary_eq = str(row[primary_eq_col] or '').strip() if primary_eq_col and row[primary_eq_col] else ''
    
    if isin and len(isin) >= 12:
        isin_count += 1
    if permid and permid != '@NA':
        permid_count += 1
    if entity_id and entity_id != '@NA':
        entity_id_count += 1
    if perm_sec and perm_sec != '@NA':
        perm_sec_count += 1
    
    if 'NVIDIA' in name.upper():
        nvidia_rows.append({
            'name': name,
            'isin': isin,
            'cusip': cusip,
            'sedol': sedol,
            'permid': permid,
            'entity_id': entity_id,
            'perm_sec_id': perm_sec,
            'primary_eq': primary_eq
        })

print(f"\nTotal rows: {rows:,}")
print(f"Has ISIN (12+ chars): {isin_count:,}")
print(f"Has Thomson PermID: {permid_count:,}")
print(f"Has EntityID: {entity_id_count:,}")
print(f"Has Perm.Sec.ID: {perm_sec_count:,}")

print(f"\nNVIDIA rows: {len(nvidia_rows)}")
for r in nvidia_rows[:5]:
    print(f"  Name: {r['name']}")
    print(f"  ISIN: {r['isin']}")
    print(f"  CUSIP: {r['cusip']}")
    print(f"  SEDOL: {r['sedol']}")
    print(f"  Thomson PermID: {r['permid']}")
    print(f"  EntityID: {r['entity_id']}")
    print(f"  Perm.Sec.ID: {r['perm_sec_id']}")
    print(f"  Primary Equity: {r['primary_eq']}")
    print("---")

wb.close()
