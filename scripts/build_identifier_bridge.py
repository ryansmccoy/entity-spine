#!/usr/bin/env python
"""
Build a comprehensive identifier bridge from FactSet files.

Links:
- FactSet Excel (public, good ISINs) 
- FactSet CSV (pub+private, 303k)
- Bloomberg pub/priv (35k with FIGI+ISIN)
- Thomson (loaded with LEI)

Output: A bridge table that can link entities across all sources.
"""

import csv
import openpyxl
from collections import defaultdict

# =============================================================================
# Step 1: Load FactSet Excel (public companies with good ISINs)
# =============================================================================
print("=" * 60)
print("Step 1: Loading FactSet Excel (public companies)")
print("=" * 60)

wb = openpyxl.load_workbook(
    r'B:\github\py-sec-edgar\data\FF_PUBLIC_COMPANIES_EXPORT_GLOBAL_IDENTIFIERS_NEW ALL.xlsx',
    read_only=True
)
ws = wb.active
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
col_map = {h: i for i, h in enumerate(headers) if h}

# Build mappings from Excel
excel_by_primary_eq = {}  # primary_equity -> record
excel_by_isin = {}        # isin -> record
excel_by_cusip = {}       # cusip -> record

for row in ws.iter_rows(min_row=2, values_only=True):
    record = {
        'name': str(row[col_map.get('Company Name')] or '').strip(),
        'isin': str(row[col_map.get('ISIN')] or '').strip(),
        'cusip': str(row[col_map.get('CUSIP')] or '').strip(),
        'sedol': str(row[col_map.get('SEDOL')] or '').strip(),
        'entity_id': str(row[col_map.get('EntityID')] or '').strip(),
        'perm_sec_id': str(row[col_map.get('Perm.Sec.ID')] or '').strip(),
        'primary_eq': str(row[col_map.get('PrimaryEquityListing')] or '').strip(),
        'thomson_permid': str(row[col_map.get('FECo InfoPermid')] or '').strip(),
    }
    
    if record['primary_eq']:
        excel_by_primary_eq[record['primary_eq']] = record
    if record['isin'] and len(record['isin']) >= 12:
        excel_by_isin[record['isin']] = record
    if record['cusip'] and len(record['cusip']) >= 8:
        excel_by_cusip[record['cusip']] = record

wb.close()

print(f"Excel records by Primary Equity: {len(excel_by_primary_eq):,}")
print(f"Excel records by ISIN: {len(excel_by_isin):,}")
print(f"Excel records by CUSIP: {len(excel_by_cusip):,}")

# =============================================================================
# Step 2: Load FactSet CSV (303k pub+private)
# =============================================================================
print("\n" + "=" * 60)
print("Step 2: Loading FactSet CSV (pub+private)")
print("=" * 60)

csv_by_primary_eq = {}
csv_rows = 0

with open(r'G:\FACTSET\screening\ff_global_identifiers.csv', 'r', encoding='latin-1') as f:
    reader = csv.DictReader(f)
    reader.fieldnames = [c.strip() for c in reader.fieldnames]
    
    for row in reader:
        csv_rows += 1
        primary_eq = row.get('Primary Equity Listing', '').strip()
        if primary_eq and primary_eq != '@NA':
            csv_by_primary_eq[primary_eq] = {
                'name': row.get('Company Name', '').strip(),
                'perm_sec_id': row.get('Perm. Sec. ID', '').strip(),
                'cusip_truncated': row.get('CUSIP', '').strip(),
                'isin_truncated': row.get('ISIN', '').strip(),
                'sedol': row.get('SEDOL', '').strip(),
            }

print(f"CSV total rows: {csv_rows:,}")
print(f"CSV records by Primary Equity: {len(csv_by_primary_eq):,}")

# =============================================================================
# Step 3: Load Bloomberg pub/priv (35k with FIGI+ISIN)
# =============================================================================
print("\n" + "=" * 60)
print("Step 3: Loading Bloomberg pub/priv")
print("=" * 60)

bb_by_isin = {}
bb_by_cusip = {}

with open(r'b:\github\py-sec-edgar\data\bb_pubpriv_20150101.csv', 'r', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    for row in reader:
        isin = row.get('ISIN', '').strip()
        figi = row.get('ID BB Global', '').strip()
        name = row.get('Name', '').strip()
        co_id = row.get('Co ID', '').strip()
        
        if isin and isin != '--' and len(isin) == 12:
            # Extract CUSIP from US ISIN
            if isin.startswith('US'):
                cusip = isin[2:11]
                bb_by_cusip[cusip] = {
                    'figi': figi,
                    'name': name,
                    'isin': isin,
                    'co_id': co_id
                }
            bb_by_isin[isin] = {
                'figi': figi,
                'name': name,
                'co_id': co_id
            }

print(f"Bloomberg records by ISIN: {len(bb_by_isin):,}")
print(f"Bloomberg records by CUSIP: {len(bb_by_cusip):,}")

# =============================================================================
# Step 4: Build unified bridge
# =============================================================================
print("\n" + "=" * 60)
print("Step 4: Building unified bridge")
print("=" * 60)

# Start with Excel records (have good ISINs) and enrich
bridge = []

for primary_eq, excel_rec in excel_by_primary_eq.items():
    entry = {
        'factset_entity_id': excel_rec['entity_id'],
        'factset_perm_sec_id': excel_rec['perm_sec_id'],
        'factset_primary_eq': primary_eq,
        'name': excel_rec['name'],
        'isin': excel_rec['isin'],
        'cusip': excel_rec['cusip'],
        'sedol': excel_rec['sedol'],
        'thomson_permid': excel_rec['thomson_permid'] if excel_rec['thomson_permid'] != '@NA' else None,
        'bloomberg_figi': None,
        'bloomberg_co_id': None,
        'in_csv_303k': False,
    }
    
    # Link to CSV (303k)
    if primary_eq in csv_by_primary_eq:
        entry['in_csv_303k'] = True
    
    # Link to Bloomberg via ISIN
    if excel_rec['isin'] in bb_by_isin:
        bb = bb_by_isin[excel_rec['isin']]
        entry['bloomberg_figi'] = bb['figi']
        entry['bloomberg_co_id'] = bb['co_id']
    elif excel_rec['cusip'] in bb_by_cusip:
        bb = bb_by_cusip[excel_rec['cusip']]
        entry['bloomberg_figi'] = bb['figi']
        entry['bloomberg_co_id'] = bb['co_id']
    
    bridge.append(entry)

print(f"Bridge entries: {len(bridge):,}")

# Count linkages
has_bb_figi = sum(1 for e in bridge if e['bloomberg_figi'])
has_thomson = sum(1 for e in bridge if e['thomson_permid'])
in_csv = sum(1 for e in bridge if e['in_csv_303k'])

print(f"  With Bloomberg FIGI: {has_bb_figi:,}")
print(f"  With Thomson PermID: {has_thomson:,}")
print(f"  In CSV 303k: {in_csv:,}")

# =============================================================================
# Step 5: Find NVIDIA in bridge
# =============================================================================
print("\n" + "=" * 60)
print("Step 5: NVIDIA in bridge")
print("=" * 60)

for entry in bridge:
    if 'NVIDIA' in entry['name'].upper() and 'Corporation' in entry['name']:
        print(f"Name: {entry['name']}")
        print(f"  ISIN: {entry['isin']}")
        print(f"  CUSIP: {entry['cusip']}")
        print(f"  FactSet Entity ID: {entry['factset_entity_id']}")
        print(f"  FactSet Perm Sec ID: {entry['factset_perm_sec_id']}")
        print(f"  Thomson PermID: {entry['thomson_permid']}")
        print(f"  Bloomberg FIGI: {entry['bloomberg_figi']}")
        print(f"  Bloomberg Co ID: {entry['bloomberg_co_id']}")
        print(f"  In CSV 303k: {entry['in_csv_303k']}")

# =============================================================================
# Step 6: Save bridge to CSV
# =============================================================================
print("\n" + "=" * 60)
print("Step 6: Saving bridge")
print("=" * 60)

output_file = r'b:\github\py-sec-edgar\entityspine\entityspine_data\identifier_bridge.csv'
with open(output_file, 'w', newline='', encoding='utf-8') as f:
    writer = csv.DictWriter(f, fieldnames=[
        'name', 'isin', 'cusip', 'sedol',
        'factset_entity_id', 'factset_perm_sec_id', 'factset_primary_eq',
        'thomson_permid', 'bloomberg_figi', 'bloomberg_co_id', 'in_csv_303k'
    ])
    writer.writeheader()
    writer.writerows(bridge)

print(f"Saved to: {output_file}")
print(f"Total bridge entries: {len(bridge):,}")
