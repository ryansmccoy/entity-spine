#!/usr/bin/env python
"""Link FactSet Excel (public) to FactSet CSV (pub+private) via Perm.Sec.ID."""

import csv
import openpyxl

# Load Excel file (public companies with good ISINs)
print("Loading Excel file (public companies)...")
wb = openpyxl.load_workbook(
    r'B:\github\py-sec-edgar\data\FF_PUBLIC_COMPANIES_EXPORT_GLOBAL_IDENTIFIERS_NEW ALL.xlsx',
    read_only=True
)
ws = wb.active

# Get headers
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
col_map = {h: i for i, h in enumerate(headers) if h}

# Build mapping: Perm.Sec.ID -> ISIN (from Excel)
excel_perm_sec_to_isin = {}
excel_isin_col = col_map.get('ISIN')
excel_perm_sec_col = col_map.get('Perm.Sec.ID')
excel_entity_col = col_map.get('EntityID')
excel_name_col = col_map.get('Company Name')
excel_primary_eq_col = col_map.get('PrimaryEquityListing')

for row in ws.iter_rows(min_row=2, values_only=True):
    isin = str(row[excel_isin_col] or '').strip() if row[excel_isin_col] else ''
    perm_sec = str(row[excel_perm_sec_col] or '').strip() if row[excel_perm_sec_col] else ''
    entity_id = str(row[excel_entity_col] or '').strip() if row[excel_entity_col] else ''
    name = str(row[excel_name_col] or '').strip() if row[excel_name_col] else ''
    primary_eq = str(row[excel_primary_eq_col] or '').strip() if row[excel_primary_eq_col] else ''
    
    if perm_sec and isin and len(isin) >= 12:
        # Normalize perm_sec (might have different suffix)
        perm_sec_base = perm_sec.rsplit('-', 1)[0] if '-' in perm_sec else perm_sec
        excel_perm_sec_to_isin[perm_sec] = {
            'isin': isin,
            'entity_id': entity_id,
            'name': name,
            'primary_eq': primary_eq
        }
        # Also store base version
        if perm_sec_base != perm_sec:
            excel_perm_sec_to_isin[perm_sec_base] = {
                'isin': isin,
                'entity_id': entity_id,
                'name': name,
                'primary_eq': primary_eq
            }

wb.close()
print(f"Excel Perm.Sec.ID -> ISIN mappings: {len(excel_perm_sec_to_isin):,}")

# Now load CSV file (pub+private, 303k rows)
print("\nLoading CSV file (pub+private)...")
csv_perm_secs = {}
with open(r'G:\FACTSET\screening\ff_global_identifiers.csv', 'r', encoding='latin-1') as f:
    reader = csv.DictReader(f)
    reader.fieldnames = [c.strip() for c in reader.fieldnames]
    
    for row in reader:
        perm_sec = row.get('Perm. Sec. ID', '').strip()
        name = row.get('Company Name', '').strip()
        isin_csv = row.get('ISIN', '').strip()
        cusip = row.get('CUSIP', '').strip()
        primary_eq = row.get('Primary Equity Listing', '').strip()
        
        if perm_sec and perm_sec != '@NA':
            perm_sec_base = perm_sec.rsplit('-', 1)[0] if '-' in perm_sec else perm_sec
            csv_perm_secs[perm_sec] = {
                'name': name,
                'isin_csv': isin_csv,
                'cusip': cusip,
                'primary_eq': primary_eq
            }
            if perm_sec_base != perm_sec:
                csv_perm_secs[perm_sec_base] = {
                    'name': name,
                    'isin_csv': isin_csv,
                    'cusip': cusip,
                    'primary_eq': primary_eq
                }

print(f"CSV Perm.Sec.ID count: {len(csv_perm_secs):,}")

# Find matches
matches = 0
enriched = []
for perm_sec, excel_data in excel_perm_sec_to_isin.items():
    if perm_sec in csv_perm_secs:
        matches += 1
        csv_data = csv_perm_secs[perm_sec]
        enriched.append({
            'perm_sec': perm_sec,
            'excel_name': excel_data['name'],
            'csv_name': csv_data['name'],
            'good_isin': excel_data['isin'],
            'csv_isin': csv_data['isin_csv'],
            'csv_cusip': csv_data['cusip'],
            'primary_eq': csv_data['primary_eq']
        })

print(f"\nMatches (can enrich CSV with good ISIN): {matches:,}")

# Show samples
print("\nSample enriched records:")
for r in enriched[:10]:
    print(f"  Perm.Sec.ID: {r['perm_sec']}")
    print(f"  Name: {r['excel_name']}")
    print(f"  Good ISIN: {r['good_isin']}")
    print(f"  CSV ISIN (truncated): {r['csv_isin']}")
    print(f"  CUSIP: {r['csv_cusip']}")
    print("---")

# Check NVIDIA
print("\nNVIDIA check:")
for perm_sec, data in excel_perm_sec_to_isin.items():
    if 'NVIDIA' in data['name'].upper() and 'Corporation' in data['name']:
        print(f"  Excel Perm.Sec.ID: {perm_sec}")
        print(f"  Good ISIN: {data['isin']}")
        if perm_sec in csv_perm_secs:
            csv = csv_perm_secs[perm_sec]
            print(f"  Found in CSV! CUSIP: {csv['cusip']}")
        else:
            print(f"  NOT in CSV - trying base...")
            base = perm_sec.rsplit('-', 1)[0]
            if base in csv_perm_secs:
                print(f"  Found base {base} in CSV!")
